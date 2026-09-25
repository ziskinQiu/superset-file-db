"""Excel parsing (.xlsx via openpyxl, .xls via xlrd) + Parquet export."""

import itertools
from datetime import date, datetime, time
from pathlib import Path

import pandas as pd
import xlrd

from .schema import FileSchema, SheetSchema, build_columns, unique_names


def list_sheets(path: str) -> list[SheetSchema]:
    names = _sheet_names(path)
    out = []
    for idx, name in enumerate(names):
        if _is_xls(path):
            row_count, column_count = _xls_dims(path, name)
        else:
            row_count, column_count = _xlsx_dims(path, name)
        out.append(
            SheetSchema(name=name, index=idx, row_count=row_count, column_count=column_count)
        )
    return out


def parse_schema(path: str, sheet: str | int) -> FileSchema:
    name, index = _resolve_sheet(path, sheet)
    it = _iter_all_rows(path, name)
    header = [str(v) if v is not None else "" for v in next(it, [])]
    names = unique_names(header)
    rows = (r for r in it if not _is_empty(r))
    columns = build_columns(names, itertools.islice(rows, 200))
    return FileSchema(columns=columns, sheets=[SheetSchema(name=name, index=index)])


def count_rows(path: str, sheet: str | int) -> int:
    return sum(1 for _ in _data_rows(path, sheet))


def read_rows(path: str, sheet: str | int, limit: int = 100, offset: int = 0) -> list[list]:
    skipped = itertools.islice(_data_rows(path, sheet), offset, None)
    return list(itertools.islice(skipped, limit))


def to_parquet(path: str, sheet: str | int, dest: str) -> str:
    name, _ = _resolve_sheet(path, sheet)
    dest_path = Path(dest)
    dest_path.parent.mkdir(parents=True, exist_ok=True)
    df = pd.read_excel(path, sheet_name=name)
    df.to_parquet(dest_path)
    return str(dest_path)


def _is_xls(path: str) -> bool:
    return path.lower().endswith(".xls")


def _sheet_names(path: str) -> list[str]:
    with pd.ExcelFile(path) as xf:
        return list(xf.sheet_names)


def _resolve_sheet(path: str, sheet: str | int) -> tuple[str, int]:
    names = _sheet_names(path)
    if isinstance(sheet, int):
        if not 0 <= sheet < len(names):
            raise ValueError(f"sheet index out of range: {sheet}")
        return names[sheet], sheet
    if sheet not in names:
        raise ValueError(f"unknown sheet: {sheet!r}")
    return sheet, names.index(sheet)


def _data_rows(path: str, sheet: str | int):
    name, _ = _resolve_sheet(path, sheet)
    it = _iter_all_rows(path, name)
    next(it, None)  # header
    return (r for r in it if not _is_empty(r))


def _iter_all_rows(path: str, sheet_name: str):
    if _is_xls(path):
        yield from _iter_rows_xls(path, sheet_name)
    else:
        yield from _iter_rows_xlsx(path, sheet_name)


def _iter_rows_xlsx(path: str, sheet_name: str):
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet_name]
        for row in ws.iter_rows(values_only=True):
            yield [_norm_cell(v) for v in row]
    finally:
        wb.close()


def _iter_rows_xls(path: str, sheet_name: str):
    book = xlrd.open_workbook(path)
    sh = book.sheet_by_name(sheet_name)
    for r in range(sh.nrows):
        yield [_norm_xlrd_cell(book, sh, r, c) for c in range(sh.ncols)]


def _xlsx_dims(path: str, sheet_name: str) -> tuple[int, int]:
    it = _iter_all_rows(path, sheet_name)
    width = len(next(it, []))
    count = 0
    for row in it:
        if _is_empty(row):
            continue
        count += 1
        width = max(width, len(row))
    return count, width


def _xls_dims(path: str, sheet_name: str) -> tuple[int, int]:
    book = xlrd.open_workbook(path)
    sh = book.sheet_by_name(sheet_name)
    return max(sh.nrows - 1, 0), sh.ncols


def _norm_cell(value):
    if value is None:
        return None
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, str):
        return value if value != "" else None
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def _norm_xlrd_cell(book, sh, row: int, col: int):
    cell = sh.cell(row, col)
    if cell.ctype == xlrd.XL_CELL_EMPTY:
        return None
    if cell.ctype == xlrd.XL_CELL_DATE:
        return xlrd.xldate_as_datetime(cell.value, book.datemode).isoformat()
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    return _norm_cell(cell.value)


def _is_empty(row: list) -> bool:
    return all(v is None or (isinstance(v, str) and v.strip() == "") for v in row)
