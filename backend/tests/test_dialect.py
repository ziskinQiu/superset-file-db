import io
import uuid

import pytest
from sqlalchemy import BigInteger, String
from sqlalchemy.engine.url import make_url
from sqlalchemy.exc import NoSuchTableError

from my_org.file_db import db as dbmod
from my_org.file_db.dialect import FileDBDialect, FileDBModule, set_engine_factory
from my_org.file_db.metadata import MetadataManager
from my_org.file_db.parsers.schema import ColumnSchema, SheetSchema
from my_org.file_db.query_engine import FileQueryEngine, view_name
from my_org.file_db.storage.local import LocalStorageBackend


def setup_function():
    dbmod.reset_db()
    dbmod.init_db("sqlite://")


def teardown_function():
    dbmod.reset_db()


def _env(tmp_path):
    storage = LocalStorageBackend(base_path=str(tmp_path / "files"))
    metadata = MetadataManager(dbmod.get_session())
    engine = FileQueryEngine(storage, metadata, parquet_dir=str(tmp_path / "parquet"))
    set_engine_factory(lambda: engine)
    return FileDBDialect(), engine, storage, metadata


def _register_csv(storage, metadata, text, columns, filename="a.csv"):
    file_id = str(uuid.uuid4())
    rel = storage.save_file(io.BytesIO(text.encode("utf-8")), "u1", file_id, filename)
    metadata.create_file_record(
        file_id=file_id,
        created_by="u1",
        filename=filename,
        file_size=len(text),
        file_type="csv",
        storage_path=rel,
        encoding="utf-8",
        delimiter=",",
    )
    metadata.create_columns(file_id, columns)
    return file_id


def _register_excel(tmp_path, storage, metadata, sheets):
    """sheets: list of (name, rows, is_selected)."""
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    for name, rows, _ in sheets:
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    src = tmp_path / "book.xlsx"
    wb.save(src)

    file_id = str(uuid.uuid4())
    with open(src, "rb") as f:
        rel = storage.save_file(f, "u1", file_id, "book.xlsx")
    metadata.create_file_record(
        file_id=file_id,
        created_by="u1",
        filename="book.xlsx",
        file_size=src.stat().st_size,
        file_type="excel",
        storage_path=rel,
        sheet_count=len(sheets),
    )
    sheet_ids = metadata.create_sheets(
        file_id, [SheetSchema(name=name, index=i) for i, (name, _, _) in enumerate(sheets)]
    )
    for sheet_id, (_, rows, _) in zip(sheet_ids, sheets):
        metadata.create_columns(
            file_id,
            [ColumnSchema(name=h, column_type="string") for h in rows[0]],
            sheet_id=sheet_id,
        )
    session = dbmod.get_session()
    for sheet, (_, _, selected) in zip(metadata.get_file(file_id).sheets, sheets):
        sheet.is_selected = selected
    session.commit()
    return file_id


def test_import_dbapi_module():
    mod = FileDBDialect.import_dbapi()
    assert callable(mod.connect)
    assert mod.paramstyle == "qmark"
    assert mod is FileDBModule


def test_create_connect_args():
    d = FileDBDialect()
    cargs, params = d.create_connect_args(make_url("filedb://"))
    assert cargs == []
    assert params == {"echo": False}

    _, params = d.create_connect_args(make_url("filedb://?echo=1"))
    assert params["echo"] is True


def test_table_names_and_columns(tmp_path):
    d, engine, storage, metadata = _env(tmp_path)
    csv_id = _register_csv(
        storage,
        metadata,
        "id,name\n1,a\n",
        [ColumnSchema(name="id", column_type="integer"), ColumnSchema(name="name", column_type="string")],
    )
    excel_id = _register_excel(
        tmp_path,
        storage,
        metadata,
        [("S0", [["id"], [1]], True), ("S1", [["id"], [2]], False)],
    )

    names = d.get_view_names(None)
    assert set(names) == {view_name(csv_id), view_name(excel_id, 0)}  # selected sheet only
    assert set(d.get_table_names(None)) == set(names)

    cols = d.get_columns(None, view_name(csv_id))
    assert [c["name"] for c in cols] == ["id", "name"]
    assert isinstance(cols[0]["type"], BigInteger)
    assert isinstance(cols[1]["type"], String)
    assert cols[0]["default"] is None

    sheet_cols = d.get_columns(None, view_name(excel_id, 0))
    assert [c["name"] for c in sheet_cols] == ["id"]

    with pytest.raises(NoSuchTableError):
        d.get_columns(None, "file_" + "0" * 32)


def test_execute_routes_to_engine(tmp_path, monkeypatch):
    d, engine, storage, metadata = _env(tmp_path)
    csv_id = _register_csv(
        storage,
        metadata,
        "id,name\n1,alice\n",
        [ColumnSchema(name="id", column_type="string"), ColumnSchema(name="name", column_type="string")],
    )
    ensured = []
    real = engine.ensure_view_by_name

    def spy(name):
        ensured.append(name)
        return real(name)

    monkeypatch.setattr(engine, "ensure_view_by_name", spy)

    conn = FileDBModule.connect()
    cur = conn.cursor()
    name = view_name(csv_id)
    cur.execute(f"SELECT id, name FROM {name}")
    assert cur.fetchall() == [("1", "alice")]
    assert ensured == [name]


def test_join_two_views(tmp_path, monkeypatch):
    d, engine, storage, metadata = _env(tmp_path)
    a_id = _register_csv(
        storage,
        metadata,
        "id,a\n1,x\n",
        [ColumnSchema(name="id", column_type="string"), ColumnSchema(name="a", column_type="string")],
        filename="a.csv",
    )
    b_id = _register_csv(
        storage,
        metadata,
        "id,b\n1,y\n",
        [ColumnSchema(name="id", column_type="string"), ColumnSchema(name="b", column_type="string")],
        filename="b.csv",
    )
    ensured = []
    real = engine.ensure_view_by_name

    def spy(name):
        ensured.append(name)
        return real(name)

    monkeypatch.setattr(engine, "ensure_view_by_name", spy)

    cur = FileDBModule.connect().cursor()
    cur.execute(
        f"SELECT t1.a, t2.b FROM {view_name(a_id)} t1 JOIN {view_name(b_id)} t2 ON t1.id = t2.id"
    )
    assert cur.fetchall() == [("x", "y")]
    assert set(ensured) == {view_name(a_id), view_name(b_id)}


def test_has_table(tmp_path):
    d, engine, storage, metadata = _env(tmp_path)
    csv_id = _register_csv(
        storage,
        metadata,
        "id\n1\n",
        [ColumnSchema(name="id", column_type="string")],
    )
    assert d.has_table(None, view_name(csv_id)) is True
    assert d.has_table(None, "file_" + "0" * 32) is False
