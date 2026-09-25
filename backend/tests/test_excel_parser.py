from datetime import datetime

import pandas as pd
import pytest
from openpyxl import Workbook

from my_org.file_db.parsers.excel_parser import (
    count_rows,
    list_sheets,
    parse_schema,
    read_rows,
    to_parquet,
)


def _make_xlsx(tmp_path):
    wb = Workbook()

    ws = wb.active
    ws.title = "销售数据"
    ws.append(["id", "amount"])
    ws.append([1, 2.5])
    ws.append([2, 3.5])
    ws.append([3, 4.5])

    ws2 = wb.create_sheet("明细")
    ws2.append(["name"])
    ws2.append(["张三"])

    ws3 = wb.create_sheet("Sheet3")
    ws3.append(["x"])
    ws3.append([10])
    ws3.append([20])

    path = tmp_path / "book.xlsx"
    wb.save(path)
    return path


def test_list_sheets(tmp_path):
    path = _make_xlsx(tmp_path)
    sheets = list_sheets(str(path))
    assert [s.name for s in sheets] == ["销售数据", "明细", "Sheet3"]
    assert [s.index for s in sheets] == [0, 1, 2]
    assert (sheets[0].row_count, sheets[0].column_count) == (3, 2)
    assert (sheets[1].row_count, sheets[1].column_count) == (1, 1)
    assert (sheets[2].row_count, sheets[2].column_count) == (2, 1)


def test_parse_schema_per_sheet(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "类型"
    ws.append(["count", "price", "label", "day"])
    ws.append([1, 2.5, "abc", "2026-09-25"])
    ws.append([2, 3.5, "def", "2026-01-01"])
    path = tmp_path / "types.xlsx"
    wb.save(path)

    fs = parse_schema(str(path), "类型")
    types = {c.name: (c.column_type, c.column_format) for c in fs.columns}
    assert types["count"] == ("integer", None)
    assert types["price"] == ("float", None)
    assert types["label"] == ("string", None)
    assert types["day"] == ("datetime", "%Y-%m-%d")

    # index-based sheet selection works too
    assert [c.name for c in parse_schema(str(path), 0).columns] == [
        "count",
        "price",
        "label",
        "day",
    ]


def test_read_rows_limit(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "big"
    ws.append(["id", "note"])
    for i in range(150):
        ws.append([i, "" if i == 5 else f"n{i}"])
    ws2 = wb.create_sheet("dates")
    ws2.append(["day"])
    ws2.append([datetime(2026, 9, 25, 10, 30)])
    path = tmp_path / "rows.xlsx"
    wb.save(path)

    rows = read_rows(str(path), "big", limit=100)
    assert len(rows) == 100
    assert rows[0] == [0, "n0"]
    assert rows[5] == [5, None]  # empty cell -> None
    assert count_rows(str(path), "big") == 150
    assert read_rows(str(path), "big", limit=10, offset=145) == [
        [145, "n145"],
        [146, "n146"],
        [147, "n147"],
        [148, "n148"],
        [149, "n149"],
    ]
    assert read_rows(str(path), "dates", limit=10) == [["2026-09-25T10:30:00"]]


def test_to_parquet_roundtrip(tmp_path):
    path = _make_xlsx(tmp_path)
    dest = tmp_path / "out" / "sub" / "s0.parquet"

    out = to_parquet(str(path), "销售数据", str(dest))
    assert out == str(dest)
    assert dest.exists()

    orig = pd.read_excel(path, sheet_name="销售数据")
    back = pd.read_parquet(dest)
    pd.testing.assert_frame_equal(orig, back)


def test_xls_smoke(tmp_path):
    xlwt = pytest.importorskip("xlwt")
    wb = xlwt.Workbook()
    ws = wb.add_sheet("数据")
    ws.write(0, 0, "id")
    ws.write(0, 1, "name")
    ws.write(1, 0, 1)
    ws.write(1, 1, "张三")
    ws.write(2, 0, 2)
    ws.write(2, 1, "李四")
    p = tmp_path / "t.xls"
    wb.save(str(p))

    sheets = list_sheets(str(p))
    assert sheets[0].name == "数据"
    assert sheets[0].row_count == 2
    assert count_rows(str(p), "数据") == 2
    rows = read_rows(str(p), "数据", limit=10)
    assert rows[0] == [1, "张三"]
    assert rows[1] == [2, "李四"]
