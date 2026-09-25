import io
import uuid
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

from my_org.file_db import db as dbmod
from my_org.file_db.metadata import MetadataManager
from my_org.file_db.parsers.schema import ColumnSchema, SheetSchema
from my_org.file_db.query_engine import FileQueryEngine, view_name
from my_org.file_db.storage.local import LocalStorageBackend


def setup_function():
    dbmod.reset_db()
    dbmod.init_db("sqlite://")


def teardown_function():
    dbmod.reset_db()


def _env(tmp_path):
    storage = LocalStorageBackend(base_path=str(tmp_path / "files"))
    metadata = MetadataManager(dbmod.get_session())
    engine = FileQueryEngine(storage, metadata, parquet_dir=str(tmp_path / "parquet"))
    return engine, storage, metadata


def _register_csv(storage, metadata, text, columns, delimiter=",", filename="a.csv"):
    file_id = str(uuid.uuid4())
    rel = storage.save_file(io.BytesIO(text.encode("utf-8")), "u1", file_id, filename)
    metadata.create_file_record(
        file_id=file_id,
        created_by="u1",
        filename=filename,
        file_size=len(text),
        file_type="csv",
        storage_path=rel,
        encoding="utf-8",
        delimiter=delimiter,
    )
    metadata.create_columns(file_id, columns)
    return file_id


def test_view_name_shape():
    fid = str(uuid.uuid4())
    flat = fid.replace("-", "")
    assert view_name(fid) == f"file_{flat}"
    assert view_name(fid, 2) == f"file_{flat}_2"
    assert view_name(fid) != view_name(str(uuid.uuid4()))
    assert view_name(fid, 0) != view_name(fid)


def test_csv_view_query(tmp_path):
    engine, storage, metadata = _env(tmp_path)
    fid = _register_csv(
        storage,
        metadata,
        "id,name\n1,alice\n2,bob\n",
        [ColumnSchema(name="id", column_type="string"), ColumnSchema(name="name", column_type="string")],
    )
    v = engine.ensure_view(fid)
    assert v == view_name(fid)
    df = engine.query(fid, f"SELECT * FROM {v} ORDER BY id")
    assert df.values.tolist() == [["1", "alice"], ["2", "bob"]]


def test_user_type_cast(tmp_path):
    engine, storage, metadata = _env(tmp_path)
    fid = _register_csv(storage, metadata, "v\n42\nx\n", [ColumnSchema(name="v", column_type="string")])
    cid = metadata.get_columns(fid)[0].id
    metadata.update_column_type(fid, cid, "integer", changed_by="u1")
    engine.invalidate(fid)

    v = engine.ensure_view(fid)
    values = engine.query(fid, f"SELECT v FROM {v}")["v"]
    assert [x for x in values if pd.notna(x)] == [42]
    assert values.isna().sum() == 1  # 'x' becomes NULL via TRY_CAST


def test_datetime_strptime(tmp_path):
    engine, storage, metadata = _env(tmp_path)
    fid = _register_csv(
        storage,
        metadata,
        "day\n25/09/2026\n01/01/2026\n",
        [ColumnSchema(name="day", column_type="string")],
    )
    cid = metadata.get_columns(fid)[0].id
    metadata.update_column_type(fid, cid, "datetime", column_format="%d/%m/%Y", changed_by="u1")
    engine.invalidate(fid)

    v = engine.ensure_view(fid)
    df = engine.query(fid, f"SELECT COUNT(*) AS n FROM {v} WHERE day >= TIMESTAMP '2026-02-01'")
    assert df["n"].tolist() == [1]  # only 25/09/2026


def test_excel_sheet_views(tmp_path, monkeypatch):
    engine, storage, metadata = _env(tmp_path)

    wb = Workbook()
    ws0 = wb.active
    ws0.title = "S0"
    ws0.append(["id"])
    ws0.append([1])
    ws0.append([2])
    ws1 = wb.create_sheet("S1")
    ws1.append(["id"])
    ws1.append([3])
    src = tmp_path / "book.xlsx"
    wb.save(src)

    fid = str(uuid.uuid4())
    with open(src, "rb") as f:
        rel = storage.save_file(f, "u1", fid, "book.xlsx")
    metadata.create_file_record(
        file_id=fid,
        created_by="u1",
        filename="book.xlsx",
        file_size=src.stat().st_size,
        file_type="excel",
        storage_path=rel,
        sheet_count=2,
    )
    sheet_ids = metadata.create_sheets(
        fid,
        [SheetSchema(name="S0", index=0), SheetSchema(name="S1", index=1)],
    )
    for sheet_id in sheet_ids:
        metadata.create_columns(fid, [ColumnSchema(name="id", column_type="string")], sheet_id=sheet_id)

    import my_org.file_db.query_engine as qe

    calls = []
    real_to_parquet = qe.to_parquet

    def counting(*args, **kwargs):
        calls.append(args)
        return real_to_parquet(*args, **kwargs)

    monkeypatch.setattr(qe, "to_parquet", counting)

    v0 = engine.ensure_view(fid, 0)
    v1 = engine.ensure_view(fid, 1)
    assert v0 == view_name(fid, 0) and v1 == view_name(fid, 1)

    assert engine.query(fid, f"SELECT id FROM {v0} ORDER BY id").values.tolist() == [[1], [2]]
    assert engine.query(fid, f"SELECT id FROM {v1}").values.tolist() == [[3]]

    engine.ensure_view(fid, 0)  # cached: no second parquet export
    assert len(calls) == 2  # one per distinct sheet


def test_preview_shape(tmp_path):
    engine, storage, metadata = _env(tmp_path)
    text = "id,name\n" + "\n".join(f"{i},n{i}" for i in range(150)) + "\n"
    fid = _register_csv(
        storage,
        metadata,
        text,
        [ColumnSchema(name="id", column_type="integer"), ColumnSchema(name="name", column_type="string")],
    )
    preview = engine.preview(fid)
    assert len(preview["rows"]) == 100
    assert preview["totalRows"] == 150
    assert [c["name"] for c in preview["columns"]] == ["id", "name"]
    assert all("type" in c for c in preview["columns"])


def test_invalidate_rebuilds(tmp_path):
    engine, storage, metadata = _env(tmp_path)
    fid = _register_csv(storage, metadata, "v\n1\n2\n", [ColumnSchema(name="v", column_type="string")])
    cid = metadata.get_columns(fid)[0].id

    v = engine.ensure_view(fid)
    metadata.update_column_type(fid, cid, "integer", changed_by="u1")
    engine.invalidate(fid)
    assert engine.ensure_view(fid) == v

    df = engine.query(fid, f"SELECT SUM(v) AS s FROM {v}")
    assert df["s"].tolist() == [3]  # numeric sum needs the new cast


def test_cross_view_join(tmp_path):
    engine, storage, metadata = _env(tmp_path)
    fid1 = _register_csv(
        storage,
        metadata,
        "id,a\n1,x\n",
        [ColumnSchema(name="id", column_type="string"), ColumnSchema(name="a", column_type="string")],
        filename="a.csv",
    )
    fid2 = _register_csv(
        storage,
        metadata,
        "id,b\n1,y\n",
        [ColumnSchema(name="id", column_type="string"), ColumnSchema(name="b", column_type="string")],
        filename="b.csv",
    )
    sql = (
        f"SELECT t1.a AS a, t2.b AS b FROM {view_name(fid1)} t1 "
        f"JOIN {view_name(fid2)} t2 ON t1.id = t2.id"
    )
    df = engine.query(fid1, sql)
    assert df.values.tolist() == [["x", "y"]]
